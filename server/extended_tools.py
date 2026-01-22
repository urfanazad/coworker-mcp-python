import os
import json
import subprocess
import httpx
from bs4 import BeautifulSoup
import xlsxwriter
from docx import Document
from fpdf import FPDF
from typing import List, Dict, Any
import time

# Transcription import
try:
    import speech_recognition as sr
    TRANSCRIPTION_SUPPORT = True
except ImportError:
    TRANSCRIPTION_SUPPORT = False

def browse_web(url: str) -> str:
    """Fetch and return text content from a URL."""
    try:
        response = httpx.get(url, follow_redirects=True, timeout=10.0)
        response.raise_for_status()
        soup = BeautifulSoup(response.text, 'html.parser')
        
        for script_or_style in soup(["script", "style"]):
            script_or_style.decompose()
            
        text = soup.get_text(separator='\n')
        lines = (line.strip() for line in text.splitlines())
        chunks = (phrase.strip() for line in lines for phrase in line.split("  "))
        text = '\n'.join(chunk for chunk in chunks if chunk)
        
        return text[:10000]
    except Exception as e:
        return f"Error browsing {url}: {str(e)}"

def create_excel(path: str, data: List[Dict[str, Any]]) -> str:
    """Create an Excel file using xlsxwriter."""
    try:
        workbook = xlsxwriter.Workbook(path)
        worksheet = workbook.add_worksheet()
        
        if not data:
            workbook.close()
            return f"Created empty Excel file at {path}"
            
        headers = list(data[0].keys())
        for col, header in enumerate(headers):
            worksheet.write(0, col, header)
            
        for row_idx, row_data in enumerate(data, start=1):
            for col_idx, header in enumerate(headers):
                worksheet.write(row_idx, col_idx, str(row_data.get(header, "")))
                
        workbook.close()
        return f"Successfully created Excel file at {path}"
    except Exception as e:
        return f"Error creating Excel: {str(e)}"

def create_word(path: str, content: str) -> str:
    """Create a Word document."""
    try:
        doc = Document()
        doc.add_paragraph(content)
        doc.save(path)
        return f"Successfully created Word file at {path}"
    except Exception as e:
        return f"Error creating Word file: {str(e)}"

def create_pdf(path: str, content: str) -> str:
    """Create a PDF document."""
    try:
        pdf = FPDF()
        pdf.add_page()
        pdf.set_font("helvetica", size=12)
        for line in content.split('\n'):
            pdf.multi_cell(0, 10, txt=line)
        pdf.output(path)
        return f"Successfully created PDF file at {path}"
    except Exception as e:
        return f"Error creating PDF: {str(e)}"

def execute_python_code(code: str) -> str:
    """Execute Python code locally and return stdout."""
    try:
        temp_file = "temp_exec.py"
        with open(temp_file, "w", encoding="utf-8") as f:
            f.write(code)
        
        py_exec = os.path.join(os.getcwd(), "venv", "Scripts", "python.exe")
        if not os.path.exists(py_exec):
            py_exec = "python"

        result = subprocess.run(
            [py_exec, temp_file],
            capture_output=True,
            text=True,
            timeout=30
        )
        
        if os.path.exists(temp_file):
            os.remove(temp_file)
        
        output = result.stdout
        if result.stderr:
            output += f"\nErrors:\n{result.stderr}"
        return output if output else "Code executed successfully (no output)."
    except subprocess.TimeoutExpired:
        if os.path.exists(temp_file):
            os.remove(temp_file)
        return "Error: Execution timed out (30s limit)"
    except Exception as e:
        return f"Error executing code: {str(e)}"

def search_audit_logs(query: str, workspace_root: str) -> str:
    audit_path = os.path.join(workspace_root, ".coworker_audit.jsonl")
    if not os.path.exists(audit_path):
        return "No audit logs found in this workspace."
    
    matches = []
    try:
        with open(audit_path, "r", encoding="utf-8") as f:
            for line in f:
                if query.lower() in line.lower():
                    matches.append(line.strip())
        
        if not matches:
            return f"No matches found for '{query}' in audit logs."
        
        return "\n".join(matches[-20:])
    except Exception as e:
        return f"Error searching logs: {str(e)}"

def search_google_drive(query: str) -> str:
    if not os.path.exists("credentials.json"):
        return ("Google Drive tool requires a 'credentials.json' file in the project root. "
                "Please download it from Google Cloud Console (OAuth 2.0 Client ID) and place it here.")
    return "Google Drive search logic initialized, but authentication is required. Run 'python auth_gdrive.py' first."

# --- MEETING ASSISTANT TOOLS ---

import ctypes
import time

def record_audio_native(path: str, duration: int) -> None:
    """Record audio using Windows MCI directly via ctypes (No PowerShell overhead)."""
    winmm = ctypes.windll.winmm
    mciSendString = winmm.mciSendStringA
    
    # helper to send commands
    def send(cmd):
        return mciSendString(cmd.encode("ascii"), None, 0, 0)

    try:
        # 1. Setup
        send("open new type waveaudio alias capture")
        send("set capture bitspersample 16")
        send("set capture samplespersec 16000")
        send("set capture channels 1")
        
        # 2. Record
        send("record capture")
        time.sleep(duration)
        
        # 3. Save
        # MCI requires forward slashes or escaped backslashes
        save_path = path.replace("\\", "/")
        send(f"save capture {save_path}")
        
    finally:
        # 4. Cleanup
        send("close capture")

def record_and_transcribe(duration: int = 10) -> str:
    """Record audio and transcribe it using Gemini 2.0 Flash (audio understanding)."""
    temp_wav = os.path.join(os.getcwd(), "meeting_temp.wav")
    t0 = time.time()
    
    try:
        # 1. Record audio
        record_audio_native(temp_wav, duration)
        t1 = time.time()
        print(f"[DEBUG] Audio Recording ({duration}s) took: {t1-t0:.2f}s")
        
        if not os.path.exists(temp_wav) or os.path.getsize(temp_wav) < 1000:
             return "Recording failed: Audio file was not created or is too small."

        # 2. Read WAV file and encode to base64
        with open(temp_wav, "rb") as audio_file:
            audio_content = audio_file.read()
            
        import base64
        audio_base64 = base64.b64encode(audio_content).decode('utf-8')
        t2 = time.time()
        print(f"[DEBUG] File Read & Base64 took: {t2-t1:.2f}s")
        
        # 3. Use Gemini to transcribe the audio
        api_key = os.getenv("GOOGLE_API_KEY")
        
        if not api_key:
            if os.path.exists(temp_wav):
                os.remove(temp_wav)
            return "No GOOGLE_API_KEY found. Speech recognition requires an API key."
        
        url = f"https://generativelanguage.googleapis.com/v1beta/models/gemini-2.0-flash-exp:generateContent?key={api_key}"
        
        payload = {
            "contents": [{
                "parts": [
                    {
                        "text": """You are an expert transcriber with AI noise cancellation.
1. Listen carefully to the audio and filter out background noise, static, and echo.
2. Detect the language automatically (supports ALL languages).
3. Transcribe exactly what was said.
4. If the audio is not in English, translate it to English in parentheses e.g. "Hola (Hello)".
5. Only output the final text. Do not add commentary."""
                    },
                    {
                        "inline_data": {
                            "mime_type": "audio/wav",
                            "data": audio_base64
                        }
                    }
                ]
            }]
        }
        
        response = httpx.post(url, json=payload, timeout=30.0)
        t3 = time.time()
        print(f"[DEBUG] Transcription API took: {t3-t2:.2f}s | Total: {t3-t0:.2f}s")
        
        response.raise_for_status()
        result = response.json()
        
        # Clean up
        if os.path.exists(temp_wav):
            os.remove(temp_wav)
        
        # Extract transcript
        if 'candidates' in result and len(result['candidates']) > 0:
            transcript = result['candidates'][0]['content']['parts'][0]['text'].strip()
            return f"Transcript ({duration}s): {transcript}"
        else:
            return "No speech detected. Please speak louder or closer to the microphone."
            
    except Exception as e:
        if os.path.exists(temp_wav):
            os.remove(temp_wav)
        return f"Meeting Assistant Error: {str(e)}"

# --- AI MEETING INSIGHTS ---

def analyze_transcript_with_ai(transcript: str) -> str:
    """Use Google Gemini REST API to generate smart meeting responses and insights."""
    try:
        # Get API key from environment
        api_key = os.getenv("GOOGLE_API_KEY")
        if not api_key:
            return "⚠️ No GOOGLE_API_KEY found. Set it in your environment to enable AI insights.\n\nGet your free key at: https://aistudio.google.com/apikey"
        
        url = f"https://generativelanguage.googleapis.com/v1beta/models/gemini-2.0-flash-exp:generateContent?key={api_key}"
        
        # Logic to adjust depth based on input length
        is_short = len(transcript.split()) < 20
        
        if is_short:
             prompt = f"""You are an elite AI Meeting Coach. 
TRANSCRIPT: "{transcript}"

The user said something short. Provide a single, brilliant "PERFECT RESPONSE" to keep the flow. 
Do not give analysis or questions unless absolutely necessary. Be extremely concise."""
        else:
            prompt = f"""You are an elite AI Meeting Coach (like Hedy AI). 
Your goal is to help me WIN this meeting. Don't just summarize. 
Analyze the dynamic, detect hidden concerns, and give me strategic advantages using the transcript below.

Transcript: "{transcript}"

Provide 3 specific sections:

1. **🧠 STRATEGIC INSIGHT**: What is really happening here? (e.g., "They are hesitant about price," "You are losing their attention," "They seem excited about feature X").
2. **❓ POWER QUESTIONS**: Give me 2 smart questions I should ask RIGHT NOW to take control or uncover deep needs.
3. **💬 PERFECT RESPONSE**: A short, persuasive line I can say to move the conversation forward.

Keep it short, punchy, and actionable. I am in the meeting right now."""

        payload = {
            "contents": [{
                "parts": [{"text": prompt}]
            }]
        }
        
        response = httpx.post(url, json=payload, timeout=30.0)
        t1 = time.time()
        print(f"[DEBUG] Gemini Analysis took: {t1-t0:.2f}s")
        
        response.raise_for_status()
        data = response.json()
        
        if 'candidates' in data and len(data['candidates']) > 0:
            text = data['candidates'][0]['content']['parts'][0]['text']
            return text
        else:
            return "AI returned no response. Please try again."
        
    except Exception as e:
        print(f"[DEBUG] Analysis Error: {e}")
        return f"AI Analysis Error: {str(e)}"
